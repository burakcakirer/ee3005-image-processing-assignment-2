"""Build DOCX and Overleaf report from a shared set of results and prose.
Author: Burak Bekir Çakırer / EE3005 Assignment 2. Uses python-docx.
"""
from __future__ import annotations
import json, csv, re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
REP=ROOT/'Report'
REPOSITORY_URL='https://github.com/burakcakirer/ee3005-image-processing-assignment-2'
COURSE_NAME='BIOMEDICAL SYSTEM DESIGN CAPSULE-S1'
MET=json.loads((ROOT/'Outputs/Results/experiment_metrics.json').read_text())
blocks=[]
def h(n,t):blocks.append(('h',n,t))
def p(t):blocks.append(('p',t))
def fig(filename,caption):blocks.append(('fig',filename,caption))
def table(title,heads,rows):blocks.append(('tab',title,heads,rows))
def formula(doc,latex):blocks.append(('math',doc,latex))
def repo_link():blocks.append(('repo',REPOSITORY_URL))
def reference(number,citation):blocks.append(('ref',number,citation))
def num(a,dec=2):return f'{a:.{dec}f}'

# The report is specific to the packaged images and measured results.
h(1,'Abstract')
p('This study evaluates eight complementary biomedical image processing tasks: local contrast enhancement, correlation and convolution, boundary padding, separable filtering, linear and nonlinear denoising, sharpening, gradient/second-order edge detection, and an integrated workflow. The experiments use two images from the supplied course archive and fixed synthetic noise inputs. Visual comparisons are supported by explicit calculations, PSNR, and documented parameter selections. The source implementations are maintained as independent project files and are intentionally not printed in this report.')
h(1,'Experimental setup and reproducibility')
p('This report covers all eight tasks in EE3005 Biomedical System Design Capsule-S1, Assignment 2. The radiograph-like chest image (image3.jpg, 800 x 752 pixels) is used for local contrast, border padding, and the integrated processing pipeline. The axial brain MRI example (image2.jpg, 617 x 617 pixels) is used for filtering, sharpening, and derivative operators. These images were supplied in the course assignment archive. Their external creators, licenses, acquisition parameters, and clinical context were not provided; therefore, no diagnostic claims are made. A visible Kenhub logo/watermark remains in the supplied MRI image.')
p('All experiments operate on explicitly saved 8-bit grayscale PNG inputs scaled to [0,1]. MATLAB source is supplied in MATLAB/run_assignment2.m; the experiments load the same fixed source and noise PNG files. A fully executed Python reference implementation generates the figures and numerical results in this report (NumPy/SciPy/scikit-image/OpenCV). The MATLAB implementation was prepared but not executed in this environment; library-specific AHE/CLAHE interpolation, boundary conventions, and rounding can change individual results. For an exact MATLAB-origin report, rerun the MATLAB script and refresh the tables and figures.')
p('Python reference settings: random generator seed 3005; Gaussian noise standard deviation 0.06 (normalized intensity); combined salt-and-pepper noise density 0.08 (equal salt and pepper fractions); independent salt-only and pepper-only densities 0.08. Prepared noisy realizations are frozen as PNG files so that method comparisons use identical contaminated pixels. Unless otherwise indicated, filtering uses mirrored-edge/half-sample symmetric extension, intensities are clipped to [0,1], and PSNR is computed relative to the corresponding clean input.')

h(1,'Task 1 - Adaptive histogram equalization and CLAHE')
p('The chest radiograph contains dark peripheral areas and relatively low-separation gray tones around internal thoracic structures; a single global gray-level mapping cannot independently tune each locality. Figure 1 compares the input with global histogram equalization (HE), a virtually unlimited local AHE configuration (8 x 8 tiles, clip limit 1.0), and two contrast-limited configurations (8 x 8, clip 0.01; 16 x 16, clip 0.03). Histogram bin count is 256 for processing; the plotted histograms use 64 bins for readability.')
fig('F01_contrast_histograms.png','Figure 1. Input, global HE, AHE, and two CLAHE configurations with corresponding intensity histograms.')
t1=MET['t1_contrast_metrics']
table('Table 1. Measured intensity standard deviation (normalized units)', ['Method','Whole image','Central ROI'],[
 [n,num(t1[key]['global_std'],4),num(t1[key]['central_roi_std'],4)] for n,key in [
 ('Original','Original'),('Global HE','HE'),('AHE 8x8','AHE'),('CLAHE 8x8 / .01','CLAHE_8x8_0.01'),('CLAHE 16x16 / .03','CLAHE_16x16_0.03')]])
p(f"For the defined central ROI (rows 190-469, columns 210-509 in zero-based Python indexing), local variation rises from {t1['Original']['central_roi_std']:.3f} to {t1['CLAHE_8x8_0.01']['central_roi_std']:.3f} under moderate CLAHE. Global HE spreads the overall histogram most strongly ({t1['HE']['global_std']:.3f} versus {t1['Original']['global_std']:.3f} originally) but also changes the dark background substantially. AHE treats each region locally, bringing out soft tissue and rib boundaries but also enhancing grain; CLAHE limits histogram-bin accumulation before redistribution to suppress that tendency. Reducing the tile size to 16 x 16 tiles while increasing clip limit to .03 changes enhancement spatially; it does not guarantee a higher global standard deviation.")
p('On visual comparison, CLAHE 8 x 8 / 0.01 is a cautious choice because anatomical contours become easier to distinguish without the most aggressive enhancement of background texture. Increased histogram spread is not, by itself, evidence of increased medical usefulness; subjective visibility and preservation of noise and artifacts must also be considered. [1, 2]')

h(1,'Task 2 - Correlation, convolution, and border effects')
p('A 3 x 3 kernel defines how neighboring intensities contribute to each filtered pixel. The assignment specifies the asymmetric horizontal derivative kernel and a 5 x 5 image matrix. At image center (row 3, column 3, one-based indexing), the extracted neighborhood is:')
formula('P = [25 35 45; 30 40 50; 35 45 55]    K = [1 0 -1; 2 0 -2; 1 0 -1]',r'P=\begin{bmatrix}25&35&45\\30&40&50\\35&45&55\end{bmatrix},\quad K=\begin{bmatrix}1&0&-1\\2&0&-2\\1&0&-1\end{bmatrix}')
p('Correlation uses K unchanged. The nine elementwise products are [25, 0, -45; 60, 0, -100; 35, 0, -55]. Their sum is (25 - 45) + (60 - 100) + (35 - 55) = -20 - 40 - 20 = -80.')
p('Convolution rotates K by 180 degrees: K_flip = [-1, 0, 1; -2, 0, 2; -1, 0, 1]. The products are [-25, 0, 45; -60, 0, 100; -35, 0, 55]. The sum is (-25 + 45) + (-60 + 100) + (-35 + 55) = 20 + 40 + 20 = +80. The sign reversal occurs because K is antisymmetric under 180-degree rotation. These calculations are also independently checked by the program and saved to task02_manual_calculations.txt.')
p('Border handling is evaluated on a fixed 510 x 455 chest crop using a normalized 3 x 3 arithmetic mean kernel. Zero adds artificial black samples; replication repeats the nearest available boundary pixel; mirror reflects adjacent interior pixels without repeating the edge pixel; circular wraps the opposite side of the crop into the kernel.')
fig('F02_padding_comparison.png','Figure 2. Four padding strategies; bottom row magnifies their upper-left output regions.')
pad=MET['t2_boundary_mae_vs_reflection']
table('Table 2. Mean absolute difference from mirror on one-pixel outer border', ['Padding rule','Boundary MAE'],[[n,num(pad[n],5)] for n in ['Zero','Replicate','Reflect','Circular']])
p(f"Relative to mirror, zero padding has a substantial one-pixel-border deviation ({pad['Zero']:.4f}), while replication produces {pad['Replicate']:.4f}; circular wrapping gives {pad['Circular']:.4f}. The inner image agrees across methods, because the 3 x 3 footprint remains within the image there. Mirror is a reasonable default for this non-periodic crop, but replication can be preferable for genuinely constant edge regions, and circular padding is appropriate for periodic measurements or tiled textures. No padding scheme is universally optimal. [1, 3]")

h(1,'Task 3 - Separable filtering and computational efficiency')
formula('K = (1/16)[1 2 1; 2 4 2; 1 2 1] = (1/4)[1;2;1] (1/4)[1 2 1]',r'K=\frac1{16}\begin{bmatrix}1&2&1\\2&4&2\\1&2&1\end{bmatrix}=\left(\frac14\begin{bmatrix}1\\2\\1\end{bmatrix}\right)\left(\frac14\begin{bmatrix}1&2&1\end{bmatrix}\right)')
p('Multiplying the column and row vectors yields the exact original 3 x 3 Gaussian kernel. The direct method applies all nine coefficients at each output pixel, while the separable method performs three horizontal and three vertical multiplications per pixel. This reduces the arithmetic work from K^2 to 2K (9 versus 6 multiplications per interior pixel, approximately 33% fewer); memory traffic and boundary overhead may influence real execution time.')
fig('F03_separable_gaussian.png','Figure 3. Direct 2D Gaussian versus successive horizontal/vertical 1D filtering on the MRI input.')
p(f"The observed maximum absolute difference is {MET['t3_max_abs_difference']:.3e}. This is floating-point rounding, not a visibly meaningful structural change. Using different border policies, clipping intermediate values, integer image types, or non-identical kernels would cause larger disagreements. Both implementations use the same mirrored boundary policy and double-precision operations. [1, 3]")

h(1,'Task 4 - Linear filtering for noise reduction')
formula('PSNR = 10 log10(1 / MSE),  MSE = mean[(clean - processed)^2]',r'\mathrm{PSNR}=10\log_{10}\left(\frac1{\mathrm{MSE}}\right),\quad\mathrm{MSE}=\frac1{MN}\sum_{x,y}(f(x,y)-g(x,y))^2')
p('The same clean MRI is retained throughout all comparisons. Gaussian noise causes distributed small-amplitude fluctuations (standard deviation 0.06), whereas 8% impulse noise replaces selected pixels with full black or white. Mean filters of size 3 x 3 and 5 x 5 and Gaussian filters of size 3 x 3 (sigma 0.8) and 7 x 7 (sigma 1.4) are applied to both noisy inputs. Intensities are compared on a common 0 to 1 range.')
fig('F04_gaussian_filtering.png','Figure 4. Gaussian-noise case: identical corrupted input, then four linear smoothing alternatives.')
fig('F04_impulse_filtering.png','Figure 5. Salt-and-pepper case: same linear filters and their reconstruction PSNR values.')
t4=MET['t4_psnr'];table('Table 3. Linear filtering PSNR (dB), higher is better', ['Noise model','Method','PSNR'],[[x['noise'].replace(' sigma=.06','').replace(' density=.08',''),x['filter'],num(x['psnr_db'])] for x in t4])
g=max((x for x in t4 if x['noise'].startswith('Gaussian')),key=lambda x:x['psnr_db'])
s=max((x for x in t4 if x['noise'].startswith('Salt')),key=lambda x:x['psnr_db'])
p(f"For the Gaussian realization, the best tested linear option is {g['filter']} with {g['psnr_db']:.2f} dB, up from {t4[0]['psnr_db']:.2f} dB in the noisy input. For the impulse realization, {s['filter']} reaches {s['psnr_db']:.2f} dB, although its impulse rejection is still much poorer than the 3 x 3 median result in Task 5. Increasing the smoothing footprint suppresses isolated fluctuations but also merges fine gyri/sulci outlines and other narrow transitions. PSNR indicates pixel-wise fidelity to this reference, not an independently established measure of diagnostic utility. [1, 3]")

h(1,'Task 5 - Median and nonlinear mean filtering')
p('For impulse corruption, a median selects the middle-ranked sample rather than averaging extreme black or white outliers. A small 3 x 3 median therefore often restores local backgrounds and edges without pulling each neighbor toward a single corrupted value. The 5 x 5 neighborhood is more aggressive and may remove legitimately narrow structures.')
fig('F05_median_impulse.png','Figure 6. Direct comparison of noisy input, arithmetic mean, and 3 x 3 / 5 x 5 median filters.')
formula('Geometric = product(x_i)^(1/n); Harmonic = n / sum(1/x_i); Contra(Q) = sum(x_i^(Q+1))/sum(x_i^Q)',r'G=\left(\prod_{i=1}^{n}x_i\right)^{1/n},\quad H=\frac{n}{\sum_i1/x_i},\quad C_Q=\frac{\sum_i x_i^{Q+1}}{\sum_i x_i^{Q}}')
p('The geometric mean is computed as exp(mean(log(max(x, 1/255)))) and the harmonic mean as 1 / mean(1/max(x, 1/255)), avoiding a logarithm or division by zero. This numerical intensity floor is itself a modeling choice and can alter deep-black pixels. The contra-harmonic filter uses Q=+1.5 or -1.5; when Q is negative, the same 1/255 floor prevents division by zero, but very dark neighborhoods can remain biased. Positive Q emphasizes brighter pixels, making it suited to pepper suppression; negative Q emphasizes darker pixels, making it suited to salt suppression.')
fig('F06_mean_variants.png','Figure 7. Geometric, harmonic, and arithmetic mean methods for the Gaussian-contaminated image.')
fig('F07_contraharmonic.png','Figure 8. Opposite Q values tested separately on 8% salt and 8% pepper corruption.')
t5=MET['t5_psnr'];table('Table 4. Nonlinear filtering and impulse-polarity experiments', ['Noise case','Method','PSNR (dB)'],[[x['noise'],x['filter'],num(x['psnr_db'])] for x in t5])
def t5v(noise,name):return next(x['psnr_db'] for x in t5 if x['noise']==noise and x['filter']==name)
p(f"Mixed salt-and-pepper noise is restored most successfully here by median 3 x 3 ({t5v('mixed_salt_pepper','Median 3x3'):.2f} dB), versus median 5 x 5 ({t5v('mixed_salt_pepper','Median 5x5'):.2f} dB) and mean 3 x 3 ({t5v('mixed_salt_pepper','Mean 3x3'):.2f} dB). For Gaussian noise, the geometric mean gives {t5v('gaussian','Geometric 3x3'):.2f} dB, whereas the harmonic mean gives {t5v('gaussian','Harmonic 3x3'):.2f} dB; the normal arithmetic mean gives {t5v('gaussian','Mean 3x3'):.2f} dB. Contra-harmonic +1.5 on pepper achieves {t5v('pepper_only','Contra 3x3 Q=+1.5'):.2f} dB, compared with {t5v('pepper_only','Contra 3x3 Q=-1.5'):.2f} dB for the wrong sign. The corresponding salt-only results favor -1.5 ({t5v('salt_only','Contra 3x3 Q=-1.5'):.2f} versus {t5v('salt_only','Contra 3x3 Q=+1.5'):.2f} dB). The salt-only and pepper-only baselines differ because the MRI has extensive dark-background pixels. This illustrates why experiments must report both the corruption model and the reference distribution.")
p('Recommendation based on these runs: median 3 x 3 for mixed impulse contamination; positive-Q contra-harmonic for pepper-only corruption; negative-Q contra-harmonic for salt-only corruption. Wrong-sign contra-harmonic filters can amplify the outlier type they should reject. More strongly weighted nonlinear methods can also introduce local haloing or intensity bias near abrupt edges.')

h(1,'Task 6 - Unsharp masking')
formula('mask = input - GaussianBlur(input); sharpened = clip(input + gamma*mask, 0, 1)',r'm(x,y)=f(x,y)-f_{\mathrm{blur}}(x,y),\quad g(x,y)=\operatorname{clip}(f(x,y)+\gamma m(x,y),0,1)')
p('A Gaussian 5 x 5 smoothing kernel with sigma 1.2 defines the low-frequency estimate. The residual acts as a high-pass component. Tested gains are 0.5, 1.0, and 2.5. The same untouched MRI is used as the source for all gain settings, rather than repeatedly sharpening an already processed image.')
fig('F08_unsharp.png','Figure 9. Clean input, blurred baseline, and three increasing unsharp gains.')
st=MET['t6_sharpening'];table('Table 5. Sharpening gain and clipping behavior', ['Gain gamma','Pixels requiring clipping (%)','Mean |change|'],[[k,num(v['clipped_percentage']),num(v['mean_absolute_change'],4)] for k,v in st.items()])
p(f"The fraction of pixels exceeding the valid range before clipping rises from {st['0.5']['clipped_percentage']:.2f}% at gamma 0.5 to {st['2.5']['clipped_percentage']:.2f}% at gamma 2.5. Gamma 1.0 is a restrained compromise: fine MRI boundaries become more conspicuous without the strongest halos and bright/dark ringing seen at high gain. Clipping counts include the large, nearly black exterior region and should not be interpreted as a clinical quality score. Conventional direct high-frequency enhancement (e.g., adding a Laplacian-derived high-pass response) is closely related, but its derivative kernel and frequency response need not match Gaussian-based unsharp masking.")

h(1,'Task 7 - Sobel, Laplacian, LoG, and DoG')
p('Sobel computes a pair of first derivatives and combines them as gradient magnitude. It highlights the strength of local intensity transitions and approximately encodes edge orientation. A four-neighbor Laplacian, [0 1 0; 1 -4 1; 0 1 0], is a signed second derivative; edges are commonly associated with its zero crossings, so displaying the absolute response is a visualization of edge-related activity, not a binary edge map. The LoG experiment applies Gaussian smoothing (7 x 7, sigma 1.2) before this Laplacian. DoG subtracts Gaussian smoothings at scales sigma=(1,2) and sigma=(2,4).')
fig('F09_edges.png','Figure 10. Response magnitudes from Sobel, Laplacian, LoG, and two DoG scales; last two panels illustrate noise sensitivity. Each response is individually normalized to its 99.5th percentile for display.')
ed=MET['t7_edge_stats'];table('Table 6. Derivative response sensitivity to added Gaussian noise', ['Measure','Clean MRI','Noisy MRI'],[
 ['Mean Sobel magnitude',num(ed['sobel_clean_mean'],4),num(ed['sobel_gaussian_noisy_mean'],4)],
 ['Mean absolute Laplacian',num(ed['laplacian_clean_mean_abs'],4),num(ed['laplacian_noisy_mean_abs'],4)]])
p(f"On the clean MRI, Sobel highlights the external head boundary, inner ventricles, and high-gradient tissue interfaces with readable continuity. The unsmoothed Laplacian also responds to very fine intensity changes and becomes conspicuously grainy after noise is added: mean absolute Laplacian response increases from {ed['laplacian_clean_mean_abs']:.4f} to {ed['laplacian_noisy_mean_abs']:.4f}. Sobel mean magnitude increases from {ed['sobel_clean_mean']:.4f} to {ed['sobel_gaussian_noisy_mean']:.4f}. LoG reduces fine-scale noise by suppressing high frequencies before differentiation. DoG approximates a scale-normalized LoG response for nearby scales (up to a sign and scale factor), while the wider (2,4) pair emphasizes broader structures and removes small details. Sobel is the most straightforward choice here for a readable structural-edge preview; LoG is useful for smoothed second-order response analysis [1, 4]. These outputs are illustrative rather than validated anatomical segmentations.")

h(1,'Task 8 - Integrated image processing pipeline')
p('The test input starts from the same low-contrast chest image used in Task 1. To simulate a difficult acquisition, independent Gaussian noise with sigma 0.03 is added and 2.5% of pixels are then replaced by an equal mixture of black and white impulses. The complete sequence is (1) median 3 x 3 to remove extreme local outliers before contrast scaling, (2) CLAHE with 8 x 8 tiles and clip limit 0.01 to enhance local contrast, (3) unsharp masking using Gaussian 5 x 5 sigma 1.1 and gamma 0.7 to recover edge crispness. The order matters: enhancing local contrast before removal of impulses would exaggerate contaminated pixels.')
fig('F10_pipeline.png','Figure 11. Input and every intermediate processing stage, ending with denoise/CLAHE/unsharp result.')
pipe=MET['t8_pipeline'];table('Table 7. End-to-end pipeline stage metrics', ['Stage','PSNR to clean (dB)','Image std','Entropy (bits)'],[[k,('Reference' if v['psnr_to_clean_db'] is None else num(v['psnr_to_clean_db'])),num(v['image_std'],4),num(v['entropy_bits'],3)] for k,v in pipe.items()])
p(f"The noisy image starts at {pipe['Noisy']['psnr_to_clean_db']:.2f} dB. Median filtering removes many impulses and improves reference fidelity to {pipe['Median']['psnr_to_clean_db']:.2f} dB. CLAHE changes intended local brightness mapping, reducing PSNR to {pipe['CLAHE']['psnr_to_clean_db']:.2f} dB while raising global gray-level standard deviation from {pipe['Median']['image_std']:.3f} to {pipe['CLAHE']['image_std']:.3f}; the final unsharp stage increases it slightly to {pipe['Final']['image_std']:.3f}. Entropy increases from {pipe['Median']['entropy_bits']:.3f} bits after denoising to {pipe['Final']['entropy_bits']:.3f} bits in the final result. The final stage has PSNR {pipe['Final']['psnr_to_clean_db']:.2f} dB, lower than the original noisy input; this is expected when deliberate tone remapping is penalized against an untouched reference. It does not by itself invalidate contrast enhancement.")
p('The final image displays clearer rib/lung boundaries in the radiograph illustration, but some fine grain and local intensity distortions remain. Limitations: a median filter may erase thin line structures; CLAHE can still amplify residual noise; unsharp masking may produce overshoot and ringing; no expert reference annotation is available. An alternate CLAHE -> median -> unsharp sequence may first amplify isolated impulses and then suppress some of the enhanced details; it should be re-evaluated rather than assumed to be superior. A conservative alternative for primarily Gaussian noise is Gaussian smoothing -> CLAHE -> restrained unsharp masking.')

h(1,'Cross-task findings and limitations')
p('The manual spatial operation shows that kernel orientation changes the sign for a directional, asymmetric filter. The separable Gaussian saves one-third of interior multiplications for the 3 x 3 case without materially changing the output. Gaussian noise responds well to small Gaussian smoothing, while sparse impulsive contamination is addressed more effectively by a 3 x 3 median filter. Contra-harmonic filtering depends crucially on the polarity of impulse noise, the Q sign, and the handling of zero-valued pixels. Edge operators are sensitive to both the differential order and image noise; scale selection and preprocessing are as important as the specific operator name. Finally, boosting contrast is a perceptual transformation and need not improve PSNR against the unenhanced image.')
p('Challenges and possible sources of disagreement include: differences between MATLAB adapthisteq and scikit-image CLAHE histogram clipping/interpolation; alternate definitions of reflect versus mirror at the edge; finite numerical precision; image watermarks or non-anatomical text; the use of fixed 8-bit noise samples; gamma/contrast changes that reduce direct PSNR; and the absence of clinical ground truth. These limitations are explicitly documented to prevent the numerical results from being overinterpreted.')

h(1,'Reproduction instructions, image sources, and deliverables')
p('Source files: Source_Images/Original_Samples/image3.jpg (chest radiograph) and image2.jpg (MRI) are included unchanged from the instructor-provided Image Samples folder; the other instructor-provided media, including two DICOM files, are preserved unmodified for completeness but are not analyzed. Grayscale derivatives and noisy images are in Source_Images/Prepared_Inputs. No externally sourced medical images or hidden datasets were used [7].')
p('Reproducibility materials are organized separately from the report in the repository. The MATLAB implementation is located in MATLAB/run_assignment2.m (requires Image Processing Toolbox); the executed Python reference is in Python_Reference/reproduce_results.py. Original and fixed prepared inputs are retained under Source_Images, and machine-readable metrics are in Outputs/Results. The repository README documents installation, execution, validation, and the distinction between MATLAB and Python outputs; see the MATLAB Image Processing Toolbox references for relevant MATLAB operator definitions [5, 6].')
p('The Report folder contains the present report in Word and PDF, an Overleaf-compatible main.tex, and its 11 figure files. No source code is reproduced within the report, including its appendices; complete implementation files are available only in the repository source directories. The original assignment instructions are included separately for reference. This arrangement preserves the report as a results-and-discussion document while allowing independent inspection of the implementation.')

h(1,'GitHub repository and source files')
p('The complete implementation and reproducibility materials are maintained separately from this report in the associated GitHub repository. The repository includes the MATLAB implementation, the executed Python reference, prepared image inputs, measurement tables, figures, and validation checks. No source-code listings are included in this document.')
repo_link()
p('Repository publication note: The link above is the intended GitHub address and will work after the project is published under the specified username and repository name. If the published repository is renamed, update this link in all report formats before submission.')

h(1,'References')
reference(1,'R. C. Gonzalez and R. E. Woods, Digital Image Processing, 4th ed. Pearson, 2018.')
reference(2,'K. Zuiderveld, "Contrast Limited Adaptive Histogram Equalization," in Graphics Gems IV, P. S. Heckbert, Ed. Academic Press, 1994, pp. 474-485.')
reference(3,'W. K. Pratt, Digital Image Processing, 4th ed. Wiley-Interscience, 2007.')
reference(4,'D. Marr and E. Hildreth, "Theory of edge detection," Proceedings of the Royal Society of London. Series B, Biological Sciences, vol. 207, no. 1167, pp. 187-217, 1980. DOI: 10.1098/rspb.1980.0020.')
reference(5,'MathWorks, "adapthisteq: Contrast-limited adaptive histogram equalization (CLAHE)," Image Processing Toolbox Documentation. https://www.mathworks.com/help/images/ref/adapthisteq.html (accessed 20 October 2026).')
reference(6,'MathWorks, "imfilter: N-D filtering of multidimensional images," Image Processing Toolbox Documentation. https://www.mathworks.com/help/images/ref/imfilter.html (accessed 20 October 2026).')
reference(7,'Course-provided Image Processing Assignment 2 instructions and accompanying image samples, EE3005 Biomedical System Design Capsule-S1, Abdullah Gul University, 2026. Unpublished course materials.')

# -----------------DOCX-----------------
def add_hyperlink(par, title, url):
    rel_id=par.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel_id)
    node=OxmlElement('w:r')
    props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'16619E');props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
    node.append(props)
    item=OxmlElement('w:t');item.text=title;node.append(item)
    link.append(node);par._p.append(link)

doc=Document()
sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)  # A4 academic report
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.70)
sec.left_margin=Inches(.70);sec.right_margin=Inches(.70)
sec.header_distance=Inches(.4);sec.footer_distance=Inches(.38)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9.5)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.15
for k,sz in [('Title',21),('Heading 1',13.7),('Heading 2',11.5)]:
    s=styles[k];s.font.name='Aptos Display';s.font.size=Pt(sz);s.font.color.rgb=RGBColor(27,53,85)
    if k!='Title':s.font.bold=True
    s.paragraph_format.space_before=Pt(13 if k!='Title' else 0);s.paragraph_format.space_after=Pt(6)

header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=header.add_run('EE3005  |  BIOMEDICAL SYSTEM DESIGN CAPSULE-S1  |  ASSIGNMENT 2')
run.font.name='Aptos';run.font.size=Pt(8);run.font.color.rgb=RGBColor(100,112,128)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Burak Bekir Çakırer  |  2211011008  |  Page ')
run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

title=doc.add_paragraph(style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.add_run('BIOMEDICAL SYSTEM DESIGN CAPSULE-S1\nEE3005 - Assignment 2')
sub=doc.add_paragraph();sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('Image Enhancement, Filtering and Edge Detection').bold=True
for k,v in [('Prepared by','Burak Bekir Çakırer'),('Student ID','2211011008'),('Course','EE3005 - BIOMEDICAL SYSTEM DESIGN CAPSULE-S1'),('Scope','Tasks 1-8; experimental methods and results')]:
    par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=par.add_run(f'{k}: ');r.bold=True;par.add_run(v)
doc.add_paragraph('Technical note: Figures and tables are generated by the executed Python reference. MATLAB source is provided for reproducing the coursework operations and requires separate execution. Image processing results are demonstrations, not clinical findings.')
front=doc.add_paragraph();front.alignment=WD_ALIGN_PARAGRAPH.CENTER
front.add_run('GitHub repository (active after publication): ').bold=True
add_hyperlink(front,REPOSITORY_URL,REPOSITORY_URL)

for block in blocks:
    if block[0]=='h':
        if block[2]=='References':doc.add_page_break()
        doc.add_heading(block[2],level=block[1])
    elif block[0]=='p':
        doc.add_paragraph(block[1])
    elif block[0]=='repo':
        par=doc.add_paragraph()
        add_hyperlink(par,block[1],block[1])
    elif block[0]=='ref':
        par=doc.add_paragraph(style='Normal')
        par.paragraph_format.left_indent=Inches(0.27)
        par.paragraph_format.first_line_indent=Inches(-0.27)
        par.add_run(f'[{block[1]}] ').bold=True
        par.add_run(block[2])
    elif block[0]=='fig':
        name,cap=block[1],block[2]
        par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next=True
        par.add_run().add_picture(str(REP/'Figures'/name),width=Inches(6.68))
        c=doc.add_paragraph(cap);c.alignment=WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after=Pt(11)
        for r in c.runs:r.italic=True;r.font.size=Pt(8)
    elif block[0]=='tab':
        title,heads,rows=block[1],block[2],block[3]
        pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True
        r=pp.add_run(title);r.bold=True;r.font.size=Pt(9)
        tb=doc.add_table(rows=1,cols=len(heads));tb.style='Light Shading Accent 1';tb.alignment=WD_TABLE_ALIGNMENT.CENTER
        for j,head in enumerate(heads):tb.rows[0].cells[j].text=str(head)
        for row in rows:
            cells=tb.add_row().cells
            for j,val in enumerate(row):cells[j].text=str(val)
        for row in tb.rows:
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after=Pt(0)
                    for r in pp.runs:r.font.size=Pt(8)
        doc.add_paragraph('')
    elif block[0]=='math':
        par=doc.add_paragraph();par.style='Normal';par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=par.add_run(block[1]);rr.font.name='Consolas';rr.font.size=Pt(8.3);rr.bold=True

docx_path=REP/'EE3005_Assignment2_Report.docx'
doc.save(docx_path)

# -----------------OVERLEAF-----------------
def texesc(s):
    specials={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(specials.get(ch,ch) for ch in str(s))
lines=[
 r'\documentclass[10pt,a4paper]{article}',r'\usepackage[utf8]{inputenc}',
 r'\usepackage[T1]{fontenc}',r'\usepackage{lmodern}',
 r'\usepackage[margin=20mm]{geometry}',r'\usepackage{graphicx}',
 r'\usepackage{amsmath,amssymb}',r'\usepackage{booktabs,longtable,array}',
 r'\usepackage{caption}',r'\usepackage{xcolor}',r'\usepackage{hyperref}',
 r'\usepackage{microtype}',r'\usepackage{fancyhdr}',r'\usepackage{float}',
 r'\hypersetup{colorlinks=true,linkcolor=blue,urlcolor=blue}',
 r'\pagestyle{fancy}',r'\fancyhf{}',
 r'\fancyhead[R]{\small EE3005 | Biomedical System Design Capsule-S1}',
 r'\fancyfoot[R]{\small Burak Bekir Çakırer / 2211011008 / \thepage}',
 r'\setlength{\headheight}{14pt}',r'\begin{document}',r'\begin{center}',
 r'{\LARGE\bfseries BIOMEDICAL SYSTEM DESIGN CAPSULE-S1}'+r'\\[7pt]',
 r'{\large EE3005 -- Assignment 2: Image Processing}'+r'\\[5pt]',
 r'{\large Image Enhancement, Filtering and Edge Detection}'+r'\\[13pt]',
 r'{\bfseries Burak Bekir Çakırer}'+r'\\',
 r'Student ID: 2211011008'+r'\\',
 r'Course: EE3005 -- BIOMEDICAL SYSTEM DESIGN CAPSULE-S1'+r'\\[8pt]',
 r'\textbf{GitHub repository (active after publication):}'+r'\\',
 r'\url{'+REPOSITORY_URL+'}',
 r'\end{center}',r'\smallskip',
 r'\noindent\textit{Technical note: Figures and tables are generated by the executed Python reference. MATLAB source is supplied and requires separate execution. Not for clinical decision making.}',
 r'\tableofcontents',r'\newpage'
]
for bl in blocks:
    kind=bl[0]
    if kind=='h':
        if bl[2]=='References':lines.append(r'\clearpage')
        lines.append(r'\section{'+texesc(bl[2])+'}')
    elif kind=='p':lines.extend([texesc(bl[1]),''])
    elif kind=='repo':lines.extend([r'\noindent\url{'+bl[1]+'}', ''])
    elif kind=='ref':lines.extend([r'\noindent\hangindent=1.5em\hangafter=1 ['+str(bl[1])+'] '+texesc(bl[2])+r'\par',''])
    elif kind=='fig':
        filename,cap=bl[1:]
        lines.extend([r'\begin{figure}[!htbp]',r'\centering',r'\includegraphics[width=\linewidth,height=0.72\textheight,keepaspectratio]{Figures/'+filename+'}',r'\caption{'+texesc(cap.split('. ',1)[1] if '. ' in cap else cap)+'}',r'\end{figure}'])
    elif kind=='tab':
        title,heads,rows=bl[1:]
        count=len(heads);coldef='|'+'|'.join(['p{'+str(round(150/count,1))+'mm}' for _ in range(count)])+'|'
        lines.extend([r'\begin{longtable}{'+coldef+'}',r'\caption{'+texesc(re.sub(r'^Table\s+\d+\.\s*','',title))+r'}\\',r'\hline', ' & '.join(r'\textbf{'+texesc(head)+'}' for head in heads)+r'\\ \hline',r'\endfirsthead',r'\hline',' & '.join(r'\textbf{'+texesc(head)+'}' for head in heads)+r'\\ \hline',r'\endhead'])
        for row in rows:lines.append(' & '.join(texesc(v) for v in row)+r'\\ \hline')
        lines.extend([r'\end{longtable}'])
    elif kind=='math':lines.append(r'\[ '+bl[2]+r' \]')
lines.append(r'\end{document}')
(REP/'main.tex').write_text('\n'.join(lines)+'\n')
print('DOCX:',docx_path)
print('LaTeX:',REP/'main.tex')
print('Count report figures:',len([b for b in blocks if b[0]=='fig']))
